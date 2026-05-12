import os
from docx import Document

def create_source_docs(source_folder, output_file):
    doc = Document()
    
    # Heading
    doc.add_heading('TÀI LIỆU MINH HỌA MÃ NGUỒN', 0)
    doc.add_paragraph('(Source Code Documentation)')
    doc.add_paragraph('Phần mềm Tích hợp Thiết bị N540X-16Z8Q2C-D, N540X-16Z4G8Q2C-D\nvới Hệ thống TDM Software')
    doc.add_paragraph('Hợp đồng: 12.2024.PONTICULUS/TRUST\nMã tài liệu: TDM-PONTICULUS-2025-DOC-3.2\nPhiên bản: v1.0')
    
    doc.add_page_break()

    # Duyệt file
    for root, dirs, files in os.walk(source_folder):
        for file in files:
            file_path = os.path.join(root, file)
            # Bỏ qua file .docx này nếu nó được tạo trong cùng thư mục
            if file == output_file:
                continue
            
            doc.add_heading(f'File: {file_path}', level=1)
            
            try:
                with open(file_path, 'r', encoding='utf-8') as f:
                    content = f.read()
                    doc.add_paragraph(content, style='Intense Quote')
            except Exception as e:
                doc.add_paragraph(f"Lỗi khi đọc file: {e}")
            
            doc.add_paragraph('\n' + '='*50 + '\n')

    doc.save(output_file)
    print(f"Đã tạo xong tài liệu tại: {output_file}")

# Thay đổi path ở đây ạ
source_path = r'C:\Users\VDC_COMPUTER\Documents\Claude\Projects\04-12.2024.PONTICULUS.TRUST\3.1_source_code'
create_source_docs(source_path, 'source_code_document.docx')
